import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_csv_report(records_data: list) -> str:
    """Generates CSV format text output."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "Record ID", "Student ID", "Student Name", "Roll Number",
        "Class", "Subject", "Date", "Timestamp", "Status",
        "Verification Method", "Face Confidence", "ID Confidence", "Override Reason"
    ])

    for row in records_data:
        writer.writerow([
            row.get("id"),
            row.get("student_id"),
            row.get("student_name"),
            row.get("roll_number"),
            row.get("class_name"),
            row.get("subject_name"),
            row.get("date"),
            row.get("timestamp"),
            row.get("status"),
            row.get("verification_method"),
            f"{row.get('face_confidence', 0)*100:.1f}%",
            f"{row.get('id_card_confidence', 0)*100:.1f}%",
            row.get("override_reason", "")
        ])

    return output.getvalue()


def generate_excel_report(records_data: list, title: str = "Attendance Report") -> io.BytesIO:
    """Generates formatted Excel workbook (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Data"

    # Title Block
    ws.merge_cells("A1:M1")
    ws["A1"] = f"ATTENDAI - {title.upper()}"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    headers = [
        "Record ID", "Student ID", "Student Name", "Roll No",
        "Class", "Subject", "Date", "Timestamp", "Status",
        "Method", "Face Conf", "ID Conf", "Notes / Override Reason"
    ]

    ws.append([]) # Blank row 2
    ws.append(headers) # Row 3

    # Style Header Row
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 25

    # Data Rows
    row_start = 4
    for idx, r in enumerate(records_data):
        row_num = row_start + idx
        status_str = str(r.get("status", "")).upper()
        
        status_fill_color = "E2E8F0"
        if status_str == "PRESENT":
            status_fill_color = "DCFCE7" # Light green
        elif status_str == "LATE":
            status_fill_color = "FEF9C3" # Light yellow
        elif status_str == "ABSENT":
            status_fill_color = "FEE2E2" # Light red

        row_data = [
            r.get("id"),
            r.get("student_id"),
            r.get("student_name"),
            r.get("roll_number"),
            r.get("class_name"),
            r.get("subject_name"),
            r.get("date"),
            r.get("timestamp"),
            status_str,
            r.get("verification_method"),
            f"{r.get('face_confidence', 0)*100:.1f}%",
            f"{r.get('id_card_confidence', 0)*100:.1f}%",
            r.get("override_reason", "")
        ]
        ws.append(row_data)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if col_num == 9: # Status column highlighting
                cell.fill = PatternFill(start_color=status_fill_color, end_color=status_fill_color, fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf_report(records_data: list, title: str = "Attendance Summary") -> io.BytesIO:
    """Generates PDF report document using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        textColor=colors.HexColor("#4F46E5"),
        spaceAfter=12
    )

    story.append(Paragraph(f"<b>AttendAI - {title}</b>", title_style))
    story.append(Paragraph(f"Generated on: 2026-09-06 | Offline Smart Attendance System", styles['Normal']))
    story.append(Spacer(1, 15))

    # Table Header & Data
    table_data = [[
        "Student ID", "Name", "Roll No", "Class", "Subject", "Status", "Confidence"
    ]]

    for r in records_data:
        status_text = f"<b>{str(r.get('status')).upper()}</b>"
        conf_text = f"{(r.get('face_confidence', 0.95))*100:.1f}%"
        table_data.append([
            str(r.get("student_id")),
            str(r.get("student_name")),
            str(r.get("roll_number")),
            str(r.get("class_name")),
            str(r.get("subject_name")),
            Paragraph(status_text, styles['Normal']),
            conf_text
        ])

    t = Table(table_data, colWidths=[80, 110, 60, 60, 100, 70, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#F8FAFC")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
    ]))

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer
